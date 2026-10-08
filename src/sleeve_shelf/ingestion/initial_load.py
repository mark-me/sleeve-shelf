"""Initial load: seeds the collection and its layout from a worked-out spreadsheet."""

import re
from collections import Counter
from dataclasses import dataclass, field, replace
from pathlib import Path

from openpyxl import load_workbook

from sleeve_shelf.domain import (
    Album,
    Artist,
    ArtistClusterAssignment,
    Cabinet,
    Cluster,
    EraBand,
    EraBandSource,
    FormatTokens,
    Placement,
    PlacementSource,
    Shelf,
    UnitType,
)
from sleeve_shelf.ingestion.formats import parse_format

LAYOUT_SHEET = "Kastindeling"
SHELVES_SHEET = "Vakoverzicht"
UNPLACED_SHEET = "Nog niet geplaatst"

LAYOUT_COLUMNS = ("Locatie", "Vak", "Cluster", "Era-band (artiest)", "Artiest", "Titel")
SHELVES_COLUMNS = ("Locatie", "Vak")
UNPLACED_COLUMNS = ("Cluster", "Artiest", "Titel")


@dataclass(slots=True)
class InitialLoad:
    """Everything one spreadsheet yields; a new upload replaces all of it."""

    cabinets: list[Cabinet] = field(default_factory=list)
    shelves: list[Shelf] = field(default_factory=list)
    clusters: list[Cluster] = field(default_factory=list)
    artists: list[Artist] = field(default_factory=list)
    assignments: list[ArtistClusterAssignment] = field(default_factory=list)
    era_bands: list[EraBand] = field(default_factory=list)
    albums: list[Album] = field(default_factory=list)
    placements: list[Placement] = field(default_factory=list)


def load_initial_layout(path: str | Path, cm_per_lp_unit: float | None = None) -> InitialLoad:
    """Read the spreadsheet at path into a fresh set of domain objects.

    The workbook counts shelf capacity in LP-units; cm_per_lp_unit turns that
    into an estimated shelf width, to be corrected by measuring later.
    """
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if LAYOUT_SHEET not in workbook.sheetnames:
            raise ValueError(f"Sheet '{LAYOUT_SHEET}' is missing from the spreadsheet")
        builder = _Builder(cm_per_lp_unit)
        if SHELVES_SHEET in workbook.sheetnames:
            for row in _rows(workbook[SHELVES_SHEET], SHELVES_COLUMNS):
                builder.add_shelf(row)
        for row in _rows(workbook[LAYOUT_SHEET], LAYOUT_COLUMNS):
            builder.add_album(row, placed=True)
        if UNPLACED_SHEET in workbook.sheetnames:
            for row in _rows(workbook[UNPLACED_SHEET], UNPLACED_COLUMNS):
                builder.add_album(row, placed=False)
    finally:
        workbook.close()
    return builder.finish()


def _rows(sheet, required: tuple[str, ...]):
    """Yield each data row as a dict keyed by header, skipping rows without a key value."""
    rows = sheet.iter_rows(values_only=True)
    header = [_text(cell) for cell in next(rows, ())]
    missing = [name for name in required if name not in header]
    if missing:
        raise ValueError(f"Sheet '{sheet.title}' is missing columns: {', '.join(missing)}")
    for values in rows:
        row = dict(zip(header, values))
        if all(_text(row[name]) for name in required):
            yield row


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _number(value) -> int | None:
    """A positive whole number, or None; the spreadsheet uses 0 for unknown."""
    return int(value) if isinstance(value, (int, float)) and value > 0 else None


def _format_tokens(format_text: str, disc_count: int | None) -> FormatTokens | None:
    if not format_text:
        return None
    tokens = parse_format(format_text).tokens
    # The workbook's own disc count wins over the one read from the format.
    return replace(tokens, disc_count=disc_count) if disc_count else tokens


def _is_confirmed_year(source: str) -> bool:
    """Whether the workbook's year source says the year is the original, not the pressing's."""
    return bool(source) and not source.startswith(("Discogs-jaar", "Onbekend"))


class _Builder:
    """Accumulates domain objects, handing out sequential ids per entity."""

    def __init__(self, cm_per_lp_unit: float | None = None) -> None:
        self.cm_per_lp_unit = cm_per_lp_unit
        self.result = InitialLoad()
        self.cabinets: dict[str, Cabinet] = {}
        self.shelves: dict[tuple[str, str], Shelf] = {}
        self.clusters: dict[str, Cluster] = {}
        self.artists: dict[str, Artist] = {}
        self.era_bands: dict[tuple[int, str], EraBand] = {}
        self.cluster_counts: dict[int, Counter[int]] = {}
        self.shelf_fill: Counter[int] = Counter()

    def add_shelf(self, row: dict) -> Shelf:
        cabinet_name, shelf_name = _text(row["Locatie"]), _text(row["Vak"])
        shelf = self.shelves.get((cabinet_name, shelf_name))
        if shelf is None:
            cabinet = self.cabinets.get(cabinet_name)
            if cabinet is None:
                cabinet = Cabinet(len(self.cabinets) + 1, cabinet_name)
                self.cabinets[cabinet_name] = cabinet
                self.result.cabinets.append(cabinet)
            shelf = Shelf(len(self.shelves) + 1, cabinet.id, shelf_name)
            self.shelves[cabinet_name, shelf_name] = shelf
            self.result.shelves.append(shelf)
        reachability = re.match(r"\d+", _text(row.get("Toegankelijkheid")))
        if reachability:
            shelf.reachability_score = int(reachability.group())
        capacity = row.get("Capaciteit")
        if shelf.width_cm is None and self.cm_per_lp_unit and isinstance(capacity, (int, float)):
            shelf.width_cm = round(capacity * self.cm_per_lp_unit, 1)
        return shelf

    def add_album(self, row: dict, *, placed: bool) -> None:
        artist = self._artist(_text(row["Artiest"]))
        start_year = _number(row.get("Jaar artiest"))
        if start_year and (artist.start_year is None or start_year < artist.start_year):
            artist.start_year = start_year
        cluster = self._cluster(_text(row["Cluster"]))
        self.cluster_counts.setdefault(artist.id, Counter())[cluster.id] += 1
        album = Album(
            id=len(self.result.albums) + 1,
            artist_id=artist.id,
            title=_text(row["Titel"]),
            format_tokens=_format_tokens(
                _text(row.get("Formaat")), _number(row.get("Aantal schijven"))
            ),
            original_release_year=_number(row.get("Sorteerjaar (origineel)")),
        )
        album.original_year_confirmed = album.original_release_year is not None and (
            _is_confirmed_year(_text(row.get("Jaarbron")))
        )
        era_label = _text(row.get("Era-band (artiest)"))
        if era_label:
            album.era_band_id = self._era_band(artist.id, era_label).id
        self.result.albums.append(album)
        if placed:
            shelf = self.add_shelf(row)
            self.result.placements.append(
                Placement(
                    UnitType.ALBUM,
                    album.id,
                    shelf.id,
                    self.shelf_fill[shelf.id],
                    PlacementSource.INITIAL_LOAD,
                )
            )
            self.shelf_fill[shelf.id] += 1

    def finish(self) -> InitialLoad:
        for artist_id, counts in self.cluster_counts.items():
            cluster_id = counts.most_common(1)[0][0]
            self.result.assignments.append(
                ArtistClusterAssignment(artist_id, cluster_id, confirmed=True)
            )
        return self.result

    def _artist(self, name: str) -> Artist:
        artist = self.artists.get(name)
        if artist is None:
            artist = Artist(len(self.artists) + 1, name)
            self.artists[name] = artist
            self.result.artists.append(artist)
        return artist

    def _cluster(self, name: str) -> Cluster:
        cluster = self.clusters.get(name)
        if cluster is None:
            cluster = Cluster(len(self.clusters) + 1, name)
            self.clusters[name] = cluster
            self.result.clusters.append(cluster)
        return cluster

    def _era_band(self, artist_id: int, label: str) -> EraBand:
        band = self.era_bands.get((artist_id, label))
        if band is None:
            position = sum(1 for key in self.era_bands if key[0] == artist_id)
            band = EraBand(
                len(self.era_bands) + 1, artist_id, label, position, EraBandSource.INITIAL_LOAD
            )
            self.era_bands[artist_id, label] = band
            self.result.era_bands.append(band)
        return band
