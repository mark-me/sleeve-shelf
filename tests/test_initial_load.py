"""Tests for the initial load from a spreadsheet."""

import pytest
from openpyxl import Workbook

from sleeve_shelf.domain import PlacementSource
from sleeve_shelf.ingestion.initial_load import load_initial_layout

LAYOUT_HEADER = [
    "Locatie",
    "Vak",
    "Cluster",
    "Era-band (artiest)",
    "Artiest",
    "Titel",
    "Formaat",
    "Sorteerjaar (origineel)",
    "Aantal schijven",
]


def _workbook(tmp_path, layout_rows, shelves=None, unplaced=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Kastindeling"
    sheet.append(LAYOUT_HEADER)
    for row in layout_rows:
        sheet.append(row)
    if shelves is not None:
        sheet = workbook.create_sheet("Vakoverzicht")
        sheet.append(["Locatie", "Vak", "Toegankelijkheid"])
        for row in shelves:
            sheet.append(row)
    if unplaced is not None:
        sheet = workbook.create_sheet("Nog niet geplaatst")
        sheet.append(["Cluster", "Artiest", "Titel", "Aantal schijven"])
        for row in unplaced:
            sheet.append(row)
    path = tmp_path / "layout.xlsx"
    workbook.save(path)
    return path


def test_rows_become_albums_placed_in_row_order(tmp_path):
    path = _workbook(
        tmp_path,
        [
            ["Kast", "a", "Jazz", "1960s", "Miles Davis", "Kind Of Blue", "LP, Album", 1959, 1],
            ["Kast", "a", "Jazz", "1970s", "Miles Davis", "Bitches Brew", "2xLP, Gat, 180", 0, 2],
            ["Kast", "b", "Rock", "1970s", "Tom Waits", "Closing Time", "LP + Box", 1973, 1],
        ],
    )

    result = load_initial_layout(path)

    assert [c.name for c in result.cabinets] == ["Kast"]
    assert [s.name for s in result.shelves] == ["a", "b"]
    assert [a.name for a in result.artists] == ["Miles Davis", "Tom Waits"]
    assert [(p.shelf_id, p.position) for p in result.placements] == [(1, 0), (1, 1), (2, 0)]
    assert all(p.source is PlacementSource.INITIAL_LOAD for p in result.placements)

    kind_of_blue, bitches_brew, closing_time = result.albums
    assert kind_of_blue.original_release_year == 1959
    assert bitches_brew.original_release_year is None
    assert bitches_brew.format_tokens.disc_count == 2
    assert bitches_brew.format_tokens.is_gatefold and bitches_brew.format_tokens.is_180_gram
    assert closing_time.format_tokens.is_compound

    assert [(b.artist_id, b.label, b.position) for b in result.era_bands] == [
        (1, "1960s", 0),
        (1, "1970s", 1),
        (2, "1970s", 0),
    ]
    assert kind_of_blue.era_band_id != bitches_brew.era_band_id


def test_artist_gets_its_most_frequent_cluster_confirmed(tmp_path):
    path = _workbook(
        tmp_path,
        [
            ["Kast", "a", "Latin", "1960s", "Various", "One", "LP", 0, 1],
            ["Kast", "a", "Misc", "1960s", "Various", "Two", "LP", 0, 1],
            ["Kast", "a", "Misc", "1960s", "Various", "Three", "LP", 0, 1],
        ],
    )

    result = load_initial_layout(path)

    (assignment,) = result.assignments
    misc = next(c for c in result.clusters if c.name == "Misc")
    assert assignment.cluster_id == misc.id
    assert assignment.confirmed


def test_shelf_overview_adds_empty_shelves_and_reachability(tmp_path):
    path = _workbook(
        tmp_path,
        [["Kast", "b", "Jazz", "1960s", "Miles Davis", "Kind Of Blue", "LP", 1959, 1]],
        shelves=[
            ["Kast", "a", "1 - makkelijkst"],
            ["Kast", "b", "2 (bovenste laag)"],
            ["TOTAAL", None, None],
        ],
    )

    result = load_initial_layout(path)

    assert [(s.name, s.reachability_score) for s in result.shelves] == [("a", 1), ("b", 2)]
    assert [c.name for c in result.cabinets] == ["Kast"]
    assert result.placements[0].shelf_id == 2


def test_unplaced_sheet_yields_albums_without_placement(tmp_path):
    path = _workbook(
        tmp_path,
        [["Kast", "a", "Jazz", "1960s", "Miles Davis", "Kind Of Blue", "LP", 1959, 1]],
        unplaced=[["Jazz", "Miles Davis", "Sketches Of Spain", 1]],
    )

    result = load_initial_layout(path)

    assert [a.title for a in result.albums] == ["Kind Of Blue", "Sketches Of Spain"]
    assert [p.album_id for p in result.placements] == [1]


def test_missing_column_is_reported(tmp_path):
    workbook = Workbook()
    workbook.active.title = "Kastindeling"
    workbook.active.append(["Locatie", "Vak", "Artiest"])
    path = tmp_path / "layout.xlsx"
    workbook.save(path)

    with pytest.raises(ValueError, match="Cluster, Era-band \\(artiest\\), Titel"):
        load_initial_layout(path)
